"""
XLSX-driven Polymarket execution bot.

Implements per-market exposure caps, maker-style buy logic, and
top-of-ask sell chasing per the provided specification.
"""
from __future__ import annotations

import argparse
import dataclasses
import logging
import os
import sys
import time
from dataclasses import dataclass, field
from enum import Enum
from typing import Dict, Iterable, List, Optional, Tuple

import openpyxl
import requests

from polymarket_client import PolymarketClient


class MarketStatus(str, Enum):
    INIT = "INIT"
    ACTIVE = "ACTIVE"
    HALTED = "HALTED"
    DONE = "DONE"


@dataclass
class BotConfig:
    poll_interval_ms: int = 500
    reprice_min_interval_ms: int = 1000
    max_open_orders_per_market: int = 2
    bid_depth_threshold_usd: float = 1000.0
    exposure_mode: str = "NET_COST_BASIS"
    log_level: str = "INFO"
    global_max_markets_enabled: int = 20
    max_total_open_orders: int = 200
    global_kill_switch_file: Optional[str] = None
    stale_book_seconds: int = 5
    api_base_url: str = "https://gamma-api.polymarket.com"


@dataclass
class MarketConfig:
    market_id: str
    side: str
    enabled: bool
    max_exposure_usd: float = 100.0
    max_entry_price: float = 1.0
    min_order_usd: float = 1.0
    notes: str = ""


@dataclass
class OrderState:
    order_id: str
    side: str
    price: float
    size: float
    last_updated_ms: int


@dataclass
class MarketState:
    config: MarketConfig
    status: MarketStatus = MarketStatus.INIT
    tick_size: Optional[float] = None
    net_shares: float = 0.0
    net_cost_basis_usd: float = 0.0
    last_reprice_buy_ms: int = 0
    last_reprice_sell_ms: int = 0
    last_seen_fill_ts: Optional[int] = None
    open_buy: Optional[OrderState] = None
    open_sell: Optional[OrderState] = None
    last_status_reason: str = ""


@dataclass
class OrderbookTop:
    bid_price: Optional[float]
    bid_size: Optional[float]
    ask_price: Optional[float]
    ask_size: Optional[float]
    timestamp: Optional[int]


def _normalize_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    if isinstance(value, str):
        return value.strip().lower() in {"true", "1", "yes", "y"}
    return False


def _safe_float(value: object, default: float) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def load_markets_from_xlsx(path: str) -> Tuple[Dict[str, MarketConfig], List[str]]:
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    header_map = {name.lower(): idx for idx, name in enumerate(headers)}
    required = {"market_id", "side", "enabled"}
    missing = [name for name in required if name not in header_map]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    configs: Dict[str, MarketConfig] = {}
    errors: List[str] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        market_id = str(row[header_map["market_id"]]).strip() if row[header_map["market_id"]] else ""
        side = str(row[header_map["side"]]).strip().upper() if row[header_map["side"]] else ""
        enabled = _normalize_bool(row[header_map["enabled"]])
        if not market_id or side not in {"YES", "NO"} or not enabled:
            if not market_id:
                errors.append("Row missing market_id; skipping")
            elif side not in {"YES", "NO"}:
                errors.append(f"Market {market_id}: invalid side {side}; skipping")
            continue

        max_exposure = _safe_float(row[header_map.get("max_exposure_usd", -1)], 100.0)
        max_entry_price = _safe_float(row[header_map.get("max_entry_price", -1)], 1.0)
        min_order_usd = _safe_float(row[header_map.get("min_order_usd", -1)], 1.0)
        notes = str(row[header_map.get("notes", -1)]) if header_map.get("notes") is not None else ""
        config = MarketConfig(
            market_id=market_id,
            side=side,
            enabled=True,
            max_exposure_usd=max_exposure,
            max_entry_price=max_entry_price,
            min_order_usd=min_order_usd,
            notes=notes,
        )

        if market_id in configs:
            existing = configs[market_id]
            if existing.side != config.side:
                existing.enabled = False
                errors.append(f"Market {market_id}: conflicting sides; disabling")
                continue
            existing.max_entry_price = min(existing.max_entry_price, config.max_entry_price)
            existing.max_exposure_usd = max(existing.max_exposure_usd, config.max_exposure_usd)
            existing.min_order_usd = max(existing.min_order_usd, config.min_order_usd)
        else:
            configs[market_id] = config

    return configs, errors


def _parse_orderbook_side(entries: Iterable[object]) -> Tuple[Optional[float], Optional[float]]:
    if not entries:
        return None, None
    first = next(iter(entries), None)
    if first is None:
        return None, None
    price = None
    size = None
    if hasattr(first, "price"):
        price = _safe_float(getattr(first, "price", None), None)  # type: ignore[arg-type]
        size = _safe_float(getattr(first, "size", None), None)  # type: ignore[arg-type]
    elif isinstance(first, dict):
        price = _safe_float(first.get("price"), None)
        size = _safe_float(first.get("size"), None)
    elif isinstance(first, (list, tuple)) and first:
        price = _safe_float(first[0], None)
        if len(first) > 1:
            size = _safe_float(first[1], None)
    return price, size


def _extract_timestamp(orderbook: object) -> Optional[int]:
    if isinstance(orderbook, dict):
        ts = orderbook.get("timestamp") or orderbook.get("ts")
        return int(ts) if ts else None
    if hasattr(orderbook, "timestamp"):
        ts = getattr(orderbook, "timestamp", None)
        return int(ts) if ts else None
    return None


def _infer_tick_size(orderbook: object) -> Optional[float]:
    candidates: List[float] = []
    if hasattr(orderbook, "bids"):
        bids = list(getattr(orderbook, "bids") or [])
        if len(bids) >= 2:
            p1, _ = _parse_orderbook_side([bids[0]])
            p2, _ = _parse_orderbook_side([bids[1]])
            if p1 is not None and p2 is not None:
                candidates.append(abs(p1 - p2))
    if hasattr(orderbook, "asks"):
        asks = list(getattr(orderbook, "asks") or [])
        if len(asks) >= 2:
            p1, _ = _parse_orderbook_side([asks[0]])
            p2, _ = _parse_orderbook_side([asks[1]])
            if p1 is not None and p2 is not None:
                candidates.append(abs(p2 - p1))
    if isinstance(orderbook, dict):
        bids = orderbook.get("bids") or []
        asks = orderbook.get("asks") or []
        if len(bids) >= 2:
            p1, _ = _parse_orderbook_side([bids[0]])
            p2, _ = _parse_orderbook_side([bids[1]])
            if p1 is not None and p2 is not None:
                candidates.append(abs(p1 - p2))
        if len(asks) >= 2:
            p1, _ = _parse_orderbook_side([asks[0]])
            p2, _ = _parse_orderbook_side([asks[1]])
            if p1 is not None and p2 is not None:
                candidates.append(abs(p2 - p1))
    candidates = [c for c in candidates if c and c > 0]
    return min(candidates) if candidates else None


def _parse_orderbook_top(orderbook: object) -> OrderbookTop:
    bid_price = bid_size = ask_price = ask_size = None
    if hasattr(orderbook, "bids"):
        bid_price, bid_size = _parse_orderbook_side(getattr(orderbook, "bids") or [])
    if hasattr(orderbook, "asks"):
        ask_price, ask_size = _parse_orderbook_side(getattr(orderbook, "asks") or [])
    if isinstance(orderbook, dict):
        if bid_price is None:
            bid_price, bid_size = _parse_orderbook_side(orderbook.get("bids") or [])
        if ask_price is None:
            ask_price, ask_size = _parse_orderbook_side(orderbook.get("asks") or [])
    timestamp = _extract_timestamp(orderbook)
    return OrderbookTop(bid_price, bid_size, ask_price, ask_size, timestamp)


def _parse_orders(raw_orders: Iterable[object]) -> List[OrderState]:
    orders: List[OrderState] = []
    now_ms = int(time.time() * 1000)
    for order in raw_orders or []:
        if isinstance(order, dict):
            order_id = str(order.get("id") or order.get("orderID") or order.get("order_id") or "")
            side = str(order.get("side") or "").upper()
            price = _safe_float(order.get("price"), 0.0)
            size = _safe_float(order.get("size") or order.get("quantity"), 0.0)
        else:
            order_id = str(getattr(order, "id", "") or getattr(order, "order_id", ""))
            side = str(getattr(order, "side", "")).upper()
            price = _safe_float(getattr(order, "price", None), 0.0)
            size = _safe_float(getattr(order, "size", None), 0.0)
        if not order_id:
            continue
        orders.append(OrderState(order_id=order_id, side=side, price=price, size=size, last_updated_ms=now_ms))
    return orders


def _parse_trades(raw_trades: Iterable[object]) -> List[dict]:
    trades: List[dict] = []
    for trade in raw_trades or []:
        if isinstance(trade, dict):
            trades.append(trade)
        else:
            trades.append({
                "side": getattr(trade, "side", None),
                "price": getattr(trade, "price", None),
                "size": getattr(trade, "size", None),
                "timestamp": getattr(trade, "timestamp", None),
            })
    return trades


def _compute_net_cost_basis(trades: List[dict], last_seen_ts: Optional[int]) -> Tuple[float, float, Optional[int]]:
    sorted_trades = sorted(
        [t for t in trades if t.get("timestamp") is not None],
        key=lambda t: int(t["timestamp"]),
    )
    net_shares = 0.0
    net_cost = 0.0
    latest_ts = last_seen_ts
    for trade in sorted_trades:
        ts = int(trade.get("timestamp"))
        latest_ts = max(latest_ts or ts, ts)
        side = str(trade.get("side") or "").upper()
        price = _safe_float(trade.get("price"), 0.0)
        size = _safe_float(trade.get("size"), 0.0)
        if price <= 0 or size <= 0:
            continue
        if side == "BUY":
            net_cost += price * size
            net_shares += size
        elif side == "SELL" and net_shares > 0:
            avg_cost = net_cost / net_shares if net_shares else 0
            sell_size = min(net_shares, size)
            net_cost -= avg_cost * sell_size
            net_shares -= sell_size
    if net_shares <= 0:
        net_shares = 0.0
        net_cost = 0.0
    return net_shares, net_cost, latest_ts


class XlsxExecutionBot:
    def __init__(self, config: BotConfig, client: PolymarketClient, markets: Dict[str, MarketConfig]):
        self.config = config
        self.client = client
        self.logger = logging.getLogger("xlsx_execution_bot")
        self.markets: Dict[str, MarketState] = {market_id: MarketState(config=cfg) for market_id, cfg in markets.items()}
        self.running = False

    def _fetch_market_metadata(self, market_id: str) -> Tuple[Optional[str], Optional[float]]:
        url = f"{self.config.api_base_url}/markets/{market_id}"
        try:
            response = requests.get(url, timeout=10)
            response.raise_for_status()
            data = response.json()
            status = data.get("state") or data.get("status")
            tick_size = _safe_float(data.get("tickSize"), None)
            return status, tick_size
        except Exception as exc:
            self.logger.warning("Market metadata fetch failed for %s: %s", market_id, exc)
        return None, None

    def _is_kill_switch_active(self) -> bool:
        if not self.config.global_kill_switch_file:
            return False
        return os.path.exists(self.config.global_kill_switch_file)

    def _ensure_market_state(self, market: MarketState) -> None:
        status, tick_size = self._fetch_market_metadata(market.config.market_id)
        if tick_size and not market.tick_size:
            market.tick_size = tick_size
        if status in {"closed", "resolved", "settled"}:
            market.status = MarketStatus.DONE
            market.last_status_reason = f"Market {status}"
            return
        if status not in {"open", "active", None}:
            market.status = MarketStatus.HALTED
            market.last_status_reason = f"Market status {status}"
            return
        if market.status in {MarketStatus.INIT, MarketStatus.HALTED}:
            market.status = MarketStatus.ACTIVE
            market.last_status_reason = "Market open"

    def _fetch_orderbook(self, market_id: str) -> Optional[OrderbookTop]:
        orderbook = self.client.get_orderbook(market_id)
        if not orderbook:
            return None
        return _parse_orderbook_top(orderbook)

    def _fetch_open_orders(self, market_id: str) -> List[OrderState]:
        raw_orders = self.client.get_open_orders(market_id)
        return _parse_orders(raw_orders)

    def _fetch_trades(self, market_id: str) -> List[dict]:
        if not hasattr(self.client, "get_trades"):
            return []
        try:
            raw_trades = self.client.get_trades(market_id)
        except Exception as exc:
            self.logger.warning("Failed to fetch trades for %s: %s", market_id, exc)
            return []
        return _parse_trades(raw_trades)

    def _cancel_order(self, order: OrderState) -> None:
        if self.client.cancel_order(order.order_id):
            self.logger.info("Cancelled order %s", order.order_id)

    def _place_order(self, market_id: str, side: str, price: float, size: float) -> Optional[OrderState]:
        response = self.client.place_limit_order(market_id, side, price, size)
        if not response:
            return None
        order_id = None
        if isinstance(response, dict):
            order_id = response.get("orderID") or response.get("id") or response.get("order_id")
        order_id = order_id or f"sim-{int(time.time() * 1000)}"
        return OrderState(order_id=str(order_id), side=side, price=price, size=size, last_updated_ms=int(time.time() * 1000))

    def _update_positions(self, market: MarketState) -> None:
        trades = self._fetch_trades(market.config.market_id)
        if not trades:
            return
        net_shares, net_cost, latest_ts = _compute_net_cost_basis(trades, market.last_seen_fill_ts)
        market.net_shares = net_shares
        market.net_cost_basis_usd = net_cost
        market.last_seen_fill_ts = latest_ts

    def _reconcile_orders(self, market: MarketState, open_orders: List[OrderState]) -> None:
        buy_orders = [o for o in open_orders if o.side == "BUY"]
        sell_orders = [o for o in open_orders if o.side == "SELL"]
        for extra in buy_orders[1:]:
            self._cancel_order(extra)
        for extra in sell_orders[1:]:
            self._cancel_order(extra)
        market.open_buy = buy_orders[0] if buy_orders else None
        market.open_sell = sell_orders[0] if sell_orders else None

    def _calculate_buy_price(self, top: OrderbookTop, market: MarketState) -> Optional[float]:
        if top.bid_price is None or top.ask_price is None:
            return None
        bid_size_usd = (top.bid_size or 0.0) * top.bid_price
        tick = market.tick_size or 0.01
        if bid_size_usd < self.config.bid_depth_threshold_usd:
            price = top.bid_price
        else:
            price = top.bid_price + tick
        if price >= top.ask_price:
            if top.ask_price <= market.config.max_entry_price:
                price = top.ask_price
            else:
                price = top.ask_price - tick
        if price <= 0:
            return None
        return round(price, 6)

    def _calculate_buy_size(self, market: MarketState, price: float) -> Optional[float]:
        remaining_exposure = market.config.max_exposure_usd - market.net_cost_basis_usd
        if remaining_exposure < market.config.min_order_usd:
            return None
        notional = remaining_exposure
        size = notional / price if price > 0 else 0
        if size <= 0:
            return None
        return size

    def _should_reprice(self, last_ms: int) -> bool:
        return int(time.time() * 1000) - last_ms >= self.config.reprice_min_interval_ms

    def _maintain_buy(self, market: MarketState, top: OrderbookTop) -> None:
        price = self._calculate_buy_price(top, market)
        if price is None:
            return
        size = self._calculate_buy_size(market, price)
        if size is None:
            if market.open_buy:
                self._cancel_order(market.open_buy)
                market.open_buy = None
            return
        if market.open_buy and abs(market.open_buy.price - price) < 1e-9:
            return
        if market.open_buy and not self._should_reprice(market.last_reprice_buy_ms):
            return
        if market.open_buy:
            self._cancel_order(market.open_buy)
        new_order = self._place_order(market.config.market_id, "BUY", price, size)
        if new_order:
            market.open_buy = new_order
            market.last_reprice_buy_ms = int(time.time() * 1000)
            self.logger.info("Buy order set %s @ %.4f (%s)", market.config.market_id, price, size)

    def _maintain_sell(self, market: MarketState, top: OrderbookTop) -> None:
        if market.net_shares <= 0 or top.ask_price is None:
            if market.open_sell:
                self._cancel_order(market.open_sell)
                market.open_sell = None
            return
        price = top.ask_price
        size = market.net_shares
        if market.open_sell and abs(market.open_sell.price - price) < 1e-9 and abs(market.open_sell.size - size) < 1e-9:
            return
        if market.open_sell and not self._should_reprice(market.last_reprice_sell_ms):
            return
        if market.open_sell:
            self._cancel_order(market.open_sell)
        new_order = self._place_order(market.config.market_id, "SELL", price, size)
        if new_order:
            market.open_sell = new_order
            market.last_reprice_sell_ms = int(time.time() * 1000)
            self.logger.info("Sell order set %s @ %.4f (%s)", market.config.market_id, price, size)

    def _handle_market(self, market: MarketState) -> None:
        self._ensure_market_state(market)
        if market.status == MarketStatus.DONE:
            if market.open_buy:
                self._cancel_order(market.open_buy)
                market.open_buy = None
            if market.open_sell:
                self._cancel_order(market.open_sell)
                market.open_sell = None
            return
        if market.status == MarketStatus.HALTED:
            if market.open_buy:
                self._cancel_order(market.open_buy)
                market.open_buy = None
            if market.open_sell:
                self._cancel_order(market.open_sell)
                market.open_sell = None
            return

        top = self._fetch_orderbook(market.config.market_id)
        if not top:
            self.logger.warning("No orderbook for %s", market.config.market_id)
            return
        if top.timestamp is not None:
            now = int(time.time())
            if now - int(top.timestamp) > self.config.stale_book_seconds:
                self.logger.warning("Stale orderbook for %s", market.config.market_id)
                return

        if market.tick_size is None:
            market.tick_size = _infer_tick_size(self.client.get_orderbook(market.config.market_id)) or 0.01

        self._update_positions(market)
        open_orders = self._fetch_open_orders(market.config.market_id)
        self._reconcile_orders(market, open_orders)
        self._maintain_buy(market, top)
        self._maintain_sell(market, top)

    def run(self) -> None:
        self.running = True
        self.logger.info("XLSX execution bot started")
        while self.running:
            if self._is_kill_switch_active():
                self.logger.error("Kill switch active; cancelling orders and stopping")
                for market in self.markets.values():
                    if market.open_buy:
                        self._cancel_order(market.open_buy)
                    if market.open_sell:
                        self._cancel_order(market.open_sell)
                self.running = False
                break

            active_markets = [m for m in self.markets.values() if m.config.enabled][: self.config.global_max_markets_enabled]
            total_orders = 0
            for market in active_markets:
                self._handle_market(market)
                total_orders += int(market.open_buy is not None) + int(market.open_sell is not None)
            if total_orders > self.config.max_total_open_orders:
                self.logger.error("Total open orders exceeded cap; stopping")
                self.running = False
                break
            time.sleep(self.config.poll_interval_ms / 1000)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run XLSX-driven Polymarket execution bot")
    parser.add_argument("--xlsx", required=True, help="Path to XLSX file")
    parser.add_argument("--host", required=True, help="Polymarket CLOB host")
    parser.add_argument("--private-key", required=True, help="Wallet private key")
    parser.add_argument("--chain-id", type=int, required=True, help="Chain ID")
    parser.add_argument("--api-key")
    parser.add_argument("--api-secret")
    parser.add_argument("--api-passphrase")
    parser.add_argument("--kill-switch-file")
    parser.add_argument("--log-level", default="INFO")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    markets, errors = load_markets_from_xlsx(args.xlsx)
    if errors:
        for error in errors:
            print(f"⚠️  {error}")
    if not markets:
        print("No enabled markets found; exiting")
        sys.exit(1)
    config = BotConfig(log_level=args.log_level, global_kill_switch_file=args.kill_switch_file)
    logging.basicConfig(level=getattr(logging, config.log_level.upper(), logging.INFO))
    client = PolymarketClient(
        host=args.host,
        private_key=args.private_key,
        chain_id=args.chain_id,
        api_key=args.api_key,
        api_secret=args.api_secret,
        api_passphrase=args.api_passphrase,
    )
    bot = XlsxExecutionBot(config, client, markets)
    bot.run()


if __name__ == "__main__":
    main()
